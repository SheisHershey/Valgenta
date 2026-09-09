"""
方向4：多格式支持工具

- convert_format: 文档格式转换（docx → markdown / txt）
- read_pdf: 读取 PDF 文档内容
"""
from pathlib import Path

import pdfplumber
from docx import Document
from langchain_core.tools import tool

from core.docx_processor import write_docx


@tool
def convert_format(input_path: str, output_path: str = "") -> str:
    """
    转换文档格式。
    支持：docx → markdown(.md)、docx → 纯文本(.txt)
    根据 output_path 扩展名自动判断目标格式。
    返回转换后保存的文件路径。
    """
    in_path = Path(input_path)
    if not in_path.exists():
        return f"文件不存在：{input_path}"

    if in_path.suffix.lower() != ".docx":
        return "目前仅支持从 docx 转出。"

    # 读取 docx 内容
    doc = Document(str(in_path))
    parts = []
    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            parts.append(text)
    content = "\n\n".join(parts)

    # 确定输出路径
    if output_path:
        out_path = Path(output_path)
    else:
        out_path = in_path.with_suffix(".md")

    out_path.parent.mkdir(parents=True, exist_ok=True)
    suffix = out_path.suffix.lower()

    if suffix == ".md":
        # 简单 markdown 转换
        md_lines = []
        for para in doc.paragraphs:
            text = para.text.strip()
            if not text:
                continue
            import re
            m = re.search(r"(?:Heading|标题)\s*(\d)", para.style.name or "", re.IGNORECASE)
            if m:
                level = int(m.group(1))
                md_lines.append(f"{'#' * level} {text}")
            else:
                md_lines.append(text)
        out_path.write_text("\n\n".join(md_lines), encoding="utf-8")
    elif suffix == ".txt":
        out_path.write_text(content, encoding="utf-8")
    else:
        return f"不支持的目标格式：{suffix}，支持 .md / .txt"

    return str(out_path.resolve())


@tool
def read_pdf(file_path: str, max_pages: int = 0) -> str:
    """
    读取 PDF 文档的文本内容。
    max_pages: 最大读取页数，0 表示读取全部。
    返回提取的文本内容。
    """
    path = Path(file_path)
    if not path.exists():
        return f"文件不存在：{file_path}"

    try:
        with pdfplumber.open(str(path)) as pdf:
            total = len(pdf.pages)
            pages_to_read = total if max_pages <= 0 else min(max_pages, total)
            parts = []
            for i in range(pages_to_read):
                page = pdf.pages[i]
                text = page.extract_text() or ""
                if text.strip():
                    parts.append(f"--- 第 {i+1} 页 ---\n{text}")
            result = "\n\n".join(parts)
            if max_pages > 0 and total > pages_to_read:
                result += f"\n\n（共 {total} 页，已读取前 {pages_to_read} 页）"
            return result if result else "PDF 中未提取到文本内容（可能是扫描件）。"
    except Exception as e:
        return f"读取 PDF 失败：{e}"
