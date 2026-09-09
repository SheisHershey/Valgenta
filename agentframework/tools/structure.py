"""
方向1：文档结构增强工具

- extract_outline: 提取文档大纲/标题层级
- format_reference: 格式化参考文献（GB/T 7714 / APA / MLA）
- generate_toc: 生成目录文本
"""
import re
from docx import Document
from langchain_core.tools import tool


@tool
def extract_outline(file_path: str) -> str:
    """
    提取 Word 文档的大纲（标题层级结构）。
    识别 Heading 1-9 样式的段落，输出层级化目录。
    """
    doc = Document(file_path)
    lines = ["# 文档大纲\n"]
    has_heading = False

    for para in doc.paragraphs:
        style_name = para.style.name if para.style else ""
        # 匹配 "Heading 1" / "标题 1" / "Heading 2" 等
        m = re.search(r"(?:Heading|标题)\s*(\d+)", style_name, re.IGNORECASE)
        if m and para.text.strip():
            level = int(m.group(1))
            indent = "  " * (level - 1)
            lines.append(f"{indent}{level}. {para.text.strip()}")
            has_heading = True

    if not has_heading:
        # 回退：用文本特征推测标题（短行 + 不以句号结尾）
        lines.append("\n（未检测到标准标题样式，以下为推测的短行标题）\n")
        for para in doc.paragraphs:
            text = para.text.strip()
            if text and len(text) <= 30 and not text.endswith(("。", ".", "！", "！", "？", "?")):
                lines.append(f"- {text}")

    return "\n".join(lines)


@tool
def format_reference(reference: str, style: str = "GB/T 7714") -> str:
    """
    格式化参考文献条目。
    style 可选：GB/T 7714（国标）、APA、MLA。
    reference 为原始文献信息，例如：
    "张三, 李四. 人工智能综述. 计算机学报, 2024, 47(3): 100-120."
    """
    from core.llm import get_llm

    prompt = (
        f"请将以下参考文献格式化为 {style} 标准格式。"
        f"只输出格式化后的文献条目，不要添加解释。\n\n"
        f"原始文献：\n{reference}"
    )
    return get_llm().invoke(prompt).content


@tool
def generate_toc(file_path: str) -> str:
    """
    为 Word 文档生成目录文本（基于标题层级）。
    返回可直接粘贴到文档开头的目录内容。
    """
    outline = extract_outline.invoke({"file_path": file_path})
    lines = outline.split("\n")
    toc = ["目  录", ""]
    for line in lines:
        if line.startswith("#"):
            continue
        if line.strip():
            toc.append(line)
    toc.append("")
    toc.append("（注：页码请根据实际文档手动填写或使用 Word 自动目录功能）")
    return "\n".join(toc)
